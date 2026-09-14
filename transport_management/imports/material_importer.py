"""Importer for TMS operational Materials and allowed Truck Types."""

from __future__ import annotations

import csv
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Iterable

import frappe
from frappe import _

from transport_management.imports.owned_truck_importer import normalize_truck_type


MATERIALS = "Materials"
CARGO_TYPES = "Cargo Types"
SUPPORTED_TRUCK_TYPES = {"TIPPER", "TANKER"}


@dataclass
class ParsedMaterialRow:
	row_number: int
	material_name: str
	truck_type: str
	active: int


@dataclass
class InvalidMaterialRow:
	row_number: int
	value: str
	reason: str


def normalize_material_name(value: Any) -> str:
	from transport_management.imports.owned_truck_importer import normalize_text

	return normalize_text(value)


def normalize_key(value: Any) -> str:
	return normalize_material_name(value).casefold()


def normalize_active(value: Any) -> int:
	text = normalize_key(value)
	if text in {"", "1", "yes", "y", "true", "active"}:
		return 1
	if text in {"0", "no", "n", "false", "inactive"}:
		return 0
	frappe.throw(_("Invalid Active value: {0}").format(value))


def parse_material_source(source: str | Path) -> dict[str, Any]:
	path = Path(source)
	extension = path.suffix.lower()
	if extension == ".xlsx":
		return parse_xlsx(path)
	if extension == ".csv":
		return parse_csv(path)
	frappe.throw(_("Unsupported file type: {0}").format(extension or _("unknown")))


def parse_xlsx(path: Path) -> dict[str, Any]:
	from openpyxl import load_workbook

	workbook = load_workbook(path, read_only=True, data_only=True)
	rows = []
	try:
		for sheet in workbook.worksheets:
			for row_number, row in enumerate(sheet.iter_rows(values_only=True), start=1):
				rows.append((row_number, row))
	finally:
		workbook.close()
	return parse_rows(rows)


def parse_csv(path: Path) -> dict[str, Any]:
	with path.open(newline="", encoding="utf-8-sig") as csvfile:
		reader = csv.reader(csvfile)
		return parse_rows((row for row in enumerate(reader, start=1)))


def parse_rows(rows: Iterable[tuple[int, Iterable[Any]]]) -> dict[str, Any]:
	source_rows = 0
	parsed: list[ParsedMaterialRow] = []
	invalid_rows: list[InvalidMaterialRow] = []
	header_map: dict[str, int] | None = None

	for row_number, row in rows:
		source_rows += 1
		values = [normalize_material_name(value) for value in row]
		if not any(values):
			continue

		if header_map is None:
			detected_header = detect_header(values)
			if detected_header:
				header_map = detected_header
				continue
			header_map = {"material": 0, "truck_type": 1, "active": 2}

		material_name, truck_type_source, active_source = extract_values(values, header_map)
		truck_type = normalize_truck_type(truck_type_source)
		error = None
		active = 1

		if not material_name:
			error = _("Material Name is required.")
		elif not truck_type:
			error = _("Allowed Truck Type is required.")
		elif truck_type not in SUPPORTED_TRUCK_TYPES:
			error = _("Unsupported Truck Type: {0}").format(truck_type)
		else:
			try:
				active = normalize_active(active_source)
			except Exception as exc:
				error = str(exc)

		if error:
			invalid_rows.append(InvalidMaterialRow(row_number=row_number, value=" | ".join(values), reason=error))
			continue

		parsed.append(
			ParsedMaterialRow(
				row_number=row_number,
				material_name=material_name,
				truck_type=truck_type,
				active=active,
			)
		)

	return {"source_rows": source_rows, "rows": parsed, "invalid_rows": invalid_rows}


def detect_header(values: list[str]) -> dict[str, int] | None:
	header = {normalize_key(value): idx for idx, value in enumerate(values) if value}
	material = first_index(header, ("material name", "material", "cargo name", "cargo type"))
	truck_type = first_index(header, ("allowed truck type", "truck type", "vehicle type"))
	active = first_index(header, ("active", "enabled"))
	if material is not None and truck_type is not None:
		return {"material": material, "truck_type": truck_type, "active": active if active is not None else 2}
	return None


def first_index(header: dict[str, int], aliases: tuple[str, ...]) -> int | None:
	for alias in aliases:
		if alias in header:
			return header[alias]
	return None


def extract_values(values: list[str], header_map: dict[str, int]) -> tuple[str, str, str]:
	material = values[header_map["material"]] if header_map["material"] < len(values) else ""
	truck_type = values[header_map["truck_type"]] if header_map["truck_type"] < len(values) else ""
	active = values[header_map["active"]] if header_map["active"] < len(values) else ""
	return material, truck_type, active


def compare_materials(rows: list[ParsedMaterialRow]) -> list[dict[str, Any]]:
	existing = get_existing_materials()
	seen_pairs: set[tuple[str, str]] = set()
	records = []

	for row in rows:
		key = normalize_key(row.material_name)
		pair = (key, row.truck_type)
		status = "NEW"
		error = ""
		existing_name = None
		existing_types: list[str] = []

		if pair in seen_pairs:
			status = "DUPLICATE_SOURCE"
			error = _("Duplicate Material/Truck Type pair in source.")
		else:
			seen_pairs.add(pair)
			matches = existing.get(key, [])
			if len(matches) > 1:
				status = "CONFLICT"
				error = _("Multiple existing Materials match this name.")
			elif matches:
				existing_name = matches[0]["name"]
				existing_types = matches[0]["allowed_truck_types"]
				if row.truck_type in existing_types:
					status = "EXISTS"
				else:
					status = "UPDATE_COMPATIBILITY"
					error = _("Existing Material needs an additional allowed Truck Type.")

		records.append(
			{
				"material_name": row.material_name,
				"truck_type": row.truck_type,
				"active": row.active,
				"existing_material": existing_name,
				"existing_allowed_truck_types": existing_types,
				"status": status,
				"error": error,
				"row_number": row.row_number,
			}
		)

	return records


def get_existing_materials() -> dict[str, list[dict[str, Any]]]:
	existing: dict[str, list[dict[str, Any]]] = {}
	for material in frappe.get_all(CARGO_TYPES, fields=["name", "cargo_name", "active"]):
		doc = frappe.get_doc(CARGO_TYPES, material.name)
		existing.setdefault(normalize_key(material.cargo_name), []).append(
			{
				"name": material.name,
				"cargo_name": material.cargo_name,
				"active": material.active,
				"allowed_truck_types": [
					row.truck_type for row in doc.get("allowed_truck_types", []) if row.truck_type
				],
			}
		)
	return existing


def dry_run_material_import(source: str | Path) -> dict[str, Any]:
	parsed = parse_material_source(source)
	records = compare_materials(parsed["rows"])
	invalid_rows = [
		{"row_number": row.row_number, "value": row.value, "reason": row.reason}
		for row in parsed["invalid_rows"]
	]
	return {
		"summary": build_summary(parsed["source_rows"], parsed["rows"], records, parsed["invalid_rows"]),
		"records": records,
		"invalid_rows": invalid_rows,
	}


def import_new_materials(source: str | Path) -> dict[str, Any]:
	result = dry_run_material_import(source)
	records = result["records"]
	summary = result["summary"]
	created = 0
	created_compatibility_rows = 0
	failed_rows = []

	if summary["conflict"]:
		result["summary"] = build_import_summary(summary, 0, 0, 0, summary["invalid"] + summary["conflict"])
		return result

	for material_group in group_new_records(records).values():
		material_name = material_group["material_name"]
		material_rows = material_group["rows"]
		try:
			doc = frappe.new_doc(CARGO_TYPES)
			doc.cargo_name = material_name
			doc.active = max(row["active"] for row in material_rows)
			for truck_type in sorted({row["truck_type"] for row in material_rows}):
				doc.append("allowed_truck_types", {"truck_type": truck_type})
			doc.insert()
			created += 1
			created_compatibility_rows += len(doc.allowed_truck_types)
			for row in material_rows:
				row["status"] = "CREATED"
				row["created_name"] = doc.name
		except Exception as exc:
			frappe.db.rollback()
			for row in material_rows:
				row["status"] = "FAILED"
				row["error"] = str(exc)
				failed_rows.append(row)
		else:
			frappe.db.commit()

	failed = len(failed_rows) + len(result["invalid_rows"])
	skipped = sum(
		1
		for row in records
		if row["status"] in {"EXISTS", "UPDATE_COMPATIBILITY", "DUPLICATE_SOURCE", "CONFLICT"}
	)
	result["summary"] = build_import_summary(summary, created, created_compatibility_rows, skipped, failed)
	result["failed_rows"] = failed_rows
	return result


def group_new_records(records: list[dict[str, Any]]) -> dict[str, dict[str, Any]]:
	grouped: dict[str, dict[str, Any]] = {}
	for row in records:
		if row["status"] == "NEW":
			key = normalize_key(row["material_name"])
			grouped.setdefault(key, {"material_name": row["material_name"], "rows": []})["rows"].append(row)
	return grouped


def build_summary(
	source_rows: int,
	rows: list[ParsedMaterialRow],
	records: list[dict[str, Any]],
	invalid_rows: list[InvalidMaterialRow],
) -> dict[str, Any]:
	return {
		"source_rows": source_rows,
		"valid_rows": len(rows),
		"unique_materials": len({normalize_key(row.material_name) for row in rows}),
		"new": count_status(records, "NEW"),
		"exists": count_status(records, "EXISTS"),
		"update_compatibility": count_status(records, "UPDATE_COMPATIBILITY"),
		"duplicate_source": count_status(records, "DUPLICATE_SOURCE"),
		"invalid": len(invalid_rows),
		"conflict": count_status(records, "CONFLICT"),
	}


def count_status(records: list[dict[str, Any]], status: str) -> int:
	return sum(1 for row in records if row["status"] == status)


def build_import_summary(
	summary: dict[str, Any],
	created: int,
	created_compatibility_rows: int,
	skipped: int,
	failed: int,
) -> dict[str, Any]:
	import_summary = summary.copy()
	import_summary.update(
		{
			"created": created,
			"created_compatibility_rows": created_compatibility_rows,
			"skipped": skipped,
			"failed": failed,
			"final_count": frappe.db.count(CARGO_TYPES),
		}
	)
	return import_summary
