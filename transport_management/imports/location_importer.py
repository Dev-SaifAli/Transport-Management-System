"""Dry-run parser for Transport Location import workbooks."""

from __future__ import annotations

import csv
import re
from collections import OrderedDict
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Iterable

import frappe
from frappe import _


DEFAULT_COUNTRY = "United Arab Emirates"
DEFAULT_LOCATION_TYPE = "Other"
USAGE_LOADING = "Loading"
USAGE_UNLOADING = "Unloading"
USAGE_BOTH = "Both"

_WHITESPACE_RE = re.compile(r"\s+")
_NON_DATA_VALUES = {"row labels", "grand total", "total", "loading location", "unloading location"}


@dataclass
class ParsedPair:
	row_number: int
	loading: str
	unloading: str
	source: str


@dataclass
class InvalidRow:
	row_number: int
	value: str
	reason: str


@dataclass
class LocationAggregate:
	location: str
	normalized_key: str
	loading_rows: set[int] = field(default_factory=set)
	unloading_rows: set[int] = field(default_factory=set)
	source_values: set[str] = field(default_factory=set)

	@property
	def usage(self) -> str:
		if self.loading_rows and self.unloading_rows:
			return USAGE_BOTH
		if self.loading_rows:
			return USAGE_LOADING
		return USAGE_UNLOADING


def normalize_location(value: Any) -> str:
	"""Normalize only whitespace and surrounding space."""
	if value is None:
		return ""
	text = str(value).strip()
	return _WHITESPACE_RE.sub(" ", text)


def normalize_key(value: str) -> str:
	return normalize_location(value).casefold()


def is_non_data_value(value: str) -> bool:
	return normalize_key(value) in _NON_DATA_VALUES


def parse_excel_locations(source: str | Path) -> dict[str, Any]:
	"""Read an .xlsx file and extract loading/unloading pairs without DB writes."""
	from openpyxl import load_workbook

	path = Path(source)
	workbook = load_workbook(path, read_only=True, data_only=True)
	rows: list[tuple[int, Iterable[Any]]] = []

	try:
		for sheet in workbook.worksheets:
			for row_number, row in enumerate(sheet.iter_rows(values_only=True), start=1):
				rows.append((row_number, row))
	finally:
		workbook.close()

	return parse_location_rows(rows)


def parse_csv_locations(source: str | Path) -> dict[str, Any]:
	"""Read a .csv file and extract loading/unloading pairs without DB writes."""
	path = Path(source)
	with path.open(newline="", encoding="utf-8-sig") as csvfile:
		reader = csv.reader(csvfile)
		return parse_location_rows((row for row in enumerate(reader, start=1)))


def parse_location_source(source: str | Path) -> dict[str, Any]:
	path = Path(source)
	extension = path.suffix.lower()
	if extension == ".xlsx":
		return parse_excel_locations(path)
	if extension == ".csv":
		return parse_csv_locations(path)
	frappe.throw(_("Unsupported file type: {0}").format(extension or _("unknown")))


def parse_location_rows(rows: Iterable[tuple[int, Iterable[Any]]]) -> dict[str, Any]:
	pairs: list[ParsedPair] = []
	invalid_rows: list[InvalidRow] = []
	source_rows = 0

	for row_number, row in rows:
		source_rows += 1
		values = [normalize_location(value) for value in row]
		values = [value for value in values if value]
		if not values:
			continue
		if len(values) == 1 and is_non_data_value(values[0]):
			continue

		pair = parse_location_pair(values)
		if pair:
			loading, unloading, source_value = pair
			if is_non_data_value(loading) or is_non_data_value(unloading):
				continue
			pairs.append(
				ParsedPair(
					row_number=row_number,
					loading=loading,
					unloading=unloading,
					source=source_value,
				)
			)
		else:
			invalid_rows.append(
				InvalidRow(
					row_number=row_number,
					value=" | ".join(values),
					reason="Could not safely map row to one loading and one unloading location.",
				)
			)

	return {"source_rows": source_rows, "pairs": pairs, "invalid_rows": invalid_rows}


def parse_location_pair(values: list[str]) -> tuple[str, str, str] | None:
	"""Return loading/unloading pair from one worksheet row when unambiguous."""
	if len(values) >= 2:
		return values[0], values[1], " | ".join(values[:2])

	value = values[0]
	if value.count(" - ") == 1:
		loading, unloading = [normalize_location(part) for part in value.split(" - ", 1)]
		if loading and unloading:
			return loading, unloading, value
	return None


def aggregate_locations(pairs: list[ParsedPair]) -> OrderedDict[str, LocationAggregate]:
	aggregates: OrderedDict[str, LocationAggregate] = OrderedDict()
	for pair in pairs:
		add_usage(aggregates, pair.loading, pair.row_number, is_loading=True)
		add_usage(aggregates, pair.unloading, pair.row_number, is_loading=False)
	return aggregates


def add_usage(
	aggregates: OrderedDict[str, LocationAggregate],
	location: str,
	row_number: int,
	*,
	is_loading: bool,
) -> None:
	key = normalize_key(location)
	if key not in aggregates:
		aggregates[key] = LocationAggregate(
			location=location,
			normalized_key=key,
			source_values={location},
		)
	aggregate = aggregates[key]
	aggregate.source_values.add(location)
	if is_loading:
		aggregate.loading_rows.add(row_number)
	else:
		aggregate.unloading_rows.add(row_number)


def get_existing_locations_by_key() -> dict[str, list[dict[str, Any]]]:
	existing: dict[str, list[dict[str, Any]]] = {}
	for row in frappe.get_all(
		"Transport Location",
		fields=["name", "location", "location_usage", "location_type", "country", "active"],
	):
		key = normalize_key(row.location or row.name)
		existing.setdefault(key, []).append(row)
	return existing


def compare_locations(
	aggregates: OrderedDict[str, LocationAggregate],
	existing_by_key: dict[str, list[dict[str, Any]]] | None = None,
) -> list[dict[str, Any]]:
	if existing_by_key is None:
		existing_by_key = get_existing_locations_by_key()

	records = []
	for aggregate in aggregates.values():
		status = "NEW"
		existing_name = None
		existing_usage = None
		existing_matches = existing_by_key.get(aggregate.normalized_key, [])

		if len(existing_matches) > 1:
			status = "CONFLICT"
		elif existing_matches:
			existing = existing_matches[0]
			existing_name = existing.name
			existing_usage = existing.location_usage
			if existing_usage == aggregate.usage:
				status = "EXISTS"
			elif existing_usage in {USAGE_LOADING, USAGE_UNLOADING, USAGE_BOTH, None, ""}:
				status = "UPDATE_USAGE"
			else:
				status = "CONFLICT"

		records.append(
			{
				"location": aggregate.location,
				"location_usage": aggregate.usage,
				"location_type": DEFAULT_LOCATION_TYPE,
				"country": DEFAULT_COUNTRY,
				"active": 1,
				"status": status,
				"existing_name": existing_name,
				"existing_usage": existing_usage,
				"loading_rows": sorted(aggregate.loading_rows),
				"unloading_rows": sorted(aggregate.unloading_rows),
				"source_values": sorted(aggregate.source_values),
			}
		)
	return records


def build_summary(source_rows: int, pairs: list[ParsedPair], records: list[dict[str, Any]], invalid_rows: list[InvalidRow]):
	summary = {
		"source_rows": source_rows,
		"valid_rows": len(pairs),
		"unique_locations": len(records),
		"loading_only": sum(1 for row in records if row["location_usage"] == USAGE_LOADING),
		"unloading_only": sum(1 for row in records if row["location_usage"] == USAGE_UNLOADING),
		"both": sum(1 for row in records if row["location_usage"] == USAGE_BOTH),
		"new": sum(1 for row in records if row["status"] == "NEW"),
		"exists": sum(1 for row in records if row["status"] == "EXISTS"),
		"update_usage": sum(1 for row in records if row["status"] == "UPDATE_USAGE"),
		"conflicts": sum(1 for row in records if row["status"] == "CONFLICT"),
		"invalid": len(invalid_rows),
	}
	return summary


def dry_run_location_import(source: str | Path) -> dict[str, Any]:
	"""Return proposed Transport Location import changes without writing to the DB."""
	parsed = parse_location_source(source)
	aggregates = aggregate_locations(parsed["pairs"])
	records = compare_locations(aggregates)
	invalid_rows = [
		{"row_number": row.row_number, "value": row.value, "reason": row.reason}
		for row in parsed["invalid_rows"]
	]
	return {
		"summary": build_summary(parsed["source_rows"], parsed["pairs"], records, parsed["invalid_rows"]),
		"records": records,
		"invalid_rows": invalid_rows,
	}


def import_new_locations(source: str | Path) -> dict[str, Any]:
	"""Create only NEW Transport Location rows from a source file."""
	result = dry_run_location_import(source)
	records = result["records"]
	summary = result["summary"]
	created = 0
	failed_rows = []

	if summary["conflicts"]:
		result["summary"] = build_import_summary(
			summary,
			created=0,
			skipped=sum(1 for row in records if row["status"] != "CONFLICT"),
			failed=summary["conflicts"] + summary["invalid"],
		)
		return result

	for row in records:
		if row["status"] != "NEW":
			continue

		try:
			doc = frappe.new_doc("Transport Location")
			doc.location = row["location"]
			doc.location_usage = row["location_usage"]
			doc.location_type = DEFAULT_LOCATION_TYPE
			doc.country = DEFAULT_COUNTRY
			doc.active = 1
			doc.insert()
			row["status"] = "CREATED"
			row["created_name"] = doc.name
			created += 1
		except Exception as exc:
			frappe.db.rollback()
			row["status"] = "FAILED"
			row["error"] = str(exc)
			failed_rows.append(row)
		else:
			frappe.db.commit()

	failed = len(failed_rows) + len(result["invalid_rows"])
	skipped = sum(1 for row in records if row["status"] in {"EXISTS", "UPDATE_USAGE", "CONFLICT"})
	result["summary"] = build_import_summary(summary, created=created, skipped=skipped, failed=failed)
	result["failed_rows"] = failed_rows
	return result


def build_import_summary(summary: dict[str, Any], *, created: int, skipped: int, failed: int) -> dict[str, Any]:
	import_summary = summary.copy()
	import_summary.update(
		{
			"created": created,
			"skipped": skipped,
			"failed": failed,
			"final_count": frappe.db.count("Transport Location"),
		}
	)
	return import_summary
