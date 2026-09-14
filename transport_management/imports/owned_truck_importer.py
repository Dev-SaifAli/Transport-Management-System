"""Importer for AL RANA owned truck master data."""

from __future__ import annotations

import csv
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Iterable

import frappe
from frappe import _


OWNED_TRUCKS = "Owned Trucks"
OWNERSHIP_TYPE = "OWN"
DEFAULT_STATUS = "Idle"
SUPPORTED_TRUCK_TYPES = {"TIPPER", "TANKER"}
TRUCK_TYPE_ALIASES = {"BULKER": "TANKER"}
SUPPORTED_EMIRATES = {"FUJ", "DXB", "RAK"}

_WHITESPACE_RE = re.compile(r"\s+")
_PLATE_RE = re.compile(r"^(\d+)\s*(?:-\s*)?([A-Za-z]+)$")
_NON_DATA_VALUES = {"", "truck", "vehicle", "plate", "license plate", "vehicle type", "truck type", "s.no", "s no", "sr no", "no"}
_PLATE_ALIASES = {"license plate", "truck", "vehicle", "plate"}
_TYPE_ALIASES = {"truck type", "vehicle type"}


@dataclass
class ParsedTruck:
	row_number: int
	source_plate: str
	normalized_plate: str
	source_truck_type: str


@dataclass
class InvalidTruckRow:
	row_number: int
	value: str
	reason: str


def normalize_text(value: Any) -> str:
	if value is None:
		return ""
	return _WHITESPACE_RE.sub(" ", str(value).strip())


def normalize_key(value: str) -> str:
	return normalize_text(value).casefold()


def normalize_truck_type(value: Any) -> str:
	truck_type = normalize_text(value).upper()
	return TRUCK_TYPE_ALIASES.get(truck_type, truck_type)


def normalize_plate(value: Any) -> str:
	text = normalize_text(value).upper()
	match = _PLATE_RE.match(text)
	if not match:
		frappe.throw(_("Invalid license plate format: {0}").format(value))
	number, emirate = match.groups()
	if emirate not in SUPPORTED_EMIRATES:
		frappe.throw(_("Unsupported emirate code: {0}").format(emirate))
	return f"{number}-{emirate}"


def try_normalize_plate(value: Any) -> tuple[str | None, str | None]:
	try:
		return normalize_plate(value), None
	except Exception as exc:
		return None, str(exc)


def parse_owned_truck_source(source: str | Path) -> dict[str, Any]:
	path = Path(source)
	extension = path.suffix.lower()
	if extension == ".xlsx":
		return parse_xlsx(path)
	if extension == ".csv":
		return parse_csv(path)
	frappe.throw(_("Unsupported file type: {0}").format(extension or _("unknown")))


def parse_xlsx(path: Path) -> dict[str, Any]:
	from openpyxl import load_workbook

	workbook = load_workbook(path, read_only=True, data_only=True)
	rows = []
	try:
		for sheet in workbook.worksheets:
			for row_number, row in enumerate(sheet.iter_rows(values_only=True), start=1):
				rows.append((row_number, row))
	finally:
		workbook.close()
	return parse_rows(rows)


def parse_csv(path: Path) -> dict[str, Any]:
	with path.open(newline="", encoding="utf-8-sig") as csvfile:
		reader = csv.reader(csvfile)
		return parse_rows((row for row in enumerate(reader, start=1)))


def parse_rows(rows: Iterable[tuple[int, Iterable[Any]]]) -> dict[str, Any]:
	source_rows = 0
	parsed: list[ParsedTruck] = []
	invalid_rows: list[InvalidTruckRow] = []
	header_map: dict[str, int] | None = None

	for row_number, row in rows:
		source_rows += 1
		values = [normalize_text(value) for value in row]
		if not any(values):
			continue

		if header_map is None:
			detected_header = detect_header(values)
			if detected_header:
				header_map = detected_header
				continue
			header_map = {"plate": 0, "type": 1}

		source_plate, source_type = extract_values(values, header_map)
		if is_non_data(source_plate) and is_non_data(source_type):
			continue

		normalized_plate, plate_error = try_normalize_plate(source_plate)
		truck_type = normalize_truck_type(source_type)
		error = None
		if plate_error:
			error = plate_error
		elif not truck_type:
			error = _("Truck Type is required.")
		elif truck_type not in SUPPORTED_TRUCK_TYPES:
			error = _("Unsupported Truck Type: {0}").format(truck_type)

		if error:
			invalid_rows.append(InvalidTruckRow(row_number=row_number, value=" | ".join(values), reason=error))
			continue

		parsed.append(
			ParsedTruck(
				row_number=row_number,
				source_plate=source_plate,
				normalized_plate=normalized_plate or "",
				source_truck_type=truck_type,
			)
		)

	return {"source_rows": source_rows, "trucks": parsed, "invalid_rows": invalid_rows}


def detect_header(values: list[str]) -> dict[str, int] | None:
	header = {normalize_key(value): idx for idx, value in enumerate(values) if value}
	plate_indexes = [header[key] for key in _PLATE_ALIASES if key in header]
	type_indexes = [header[key] for key in _TYPE_ALIASES if key in header]
	if plate_indexes and type_indexes:
		return {"plate": plate_indexes[0], "type": type_indexes[0]}
	return None


def extract_values(values: list[str], header_map: dict[str, int]) -> tuple[str, str]:
	plate = values[header_map["plate"]] if header_map["plate"] < len(values) else ""
	truck_type = values[header_map["type"]] if header_map["type"] < len(values) else ""
	return plate, truck_type


def is_non_data(value: str) -> bool:
	return normalize_key(value) in _NON_DATA_VALUES


def get_existing_trucks_by_plate() -> dict[str, list[dict[str, Any]]]:
	existing: dict[str, list[dict[str, Any]]] = {}
	for row in frappe.get_all(
		"Truck",
		fields=["name", "truck_number", "license_plate", "vehicle_type", "ownership_type"],
	):
		for value in (row.license_plate, row.truck_number, row.name):
			normalized, _error = try_normalize_plate(value)
			if normalized:
				existing.setdefault(normalized.casefold(), []).append(row)
				break
	return existing


def compare_trucks(trucks: list[ParsedTruck]) -> list[dict[str, Any]]:
	existing_by_plate = get_existing_trucks_by_plate()
	seen: set[str] = set()
	records = []

	for truck in trucks:
		key = truck.normalized_plate.casefold()
		status = "NEW"
		error = ""
		existing_name = None
		existing_type = None

		if key in seen:
			status = "DUPLICATE_SOURCE"
			error = _("Duplicate plate in source after normalization.")
		else:
			seen.add(key)
			matches = existing_by_plate.get(key, [])
			if len(matches) > 1:
				status = "CONFLICT"
				error = _("Multiple existing Trucks match this normalized plate.")
			elif matches:
				existing = matches[0]
				existing_name = existing.name
				existing_type = existing.vehicle_type
				if existing.ownership_type and existing.ownership_type != OWNERSHIP_TYPE:
					status = "OWNERSHIP_CONFLICT"
					error = _("Existing Truck is not marked as OWN.")
				elif existing.vehicle_type and normalize_truck_type(existing.vehicle_type) != truck.source_truck_type:
					status = "UPDATE_TYPE"
					error = _("Existing Truck Type differs from source; automatic update is deferred.")
				else:
					status = "EXISTS"

		records.append(
			{
				"source_plate": truck.source_plate,
				"normalized_plate": truck.normalized_plate,
				"source_truck_type": truck.source_truck_type,
				"existing_truck": existing_name,
				"existing_type": existing_type,
				"status": status,
				"error": error,
				"row_number": truck.row_number,
			}
		)
	return records


def build_summary(source_rows: int, trucks: list[ParsedTruck], records: list[dict[str, Any]], invalid_rows: list[InvalidTruckRow]):
	return {
		"source_rows": source_rows,
		"valid_rows": len(trucks),
		"unique_trucks": len({truck.normalized_plate.casefold() for truck in trucks}),
		"new": count_status(records, "NEW"),
		"exists": count_status(records, "EXISTS"),
		"update_type": count_status(records, "UPDATE_TYPE"),
		"ownership_conflict": count_status(records, "OWNERSHIP_CONFLICT"),
		"duplicate_source": count_status(records, "DUPLICATE_SOURCE"),
		"invalid": len(invalid_rows),
		"conflict": count_status(records, "CONFLICT"),
	}


def count_status(records: list[dict[str, Any]], status: str) -> int:
	return sum(1 for row in records if row["status"] == status)


def dry_run_owned_truck_import(source: str | Path) -> dict[str, Any]:
	parsed = parse_owned_truck_source(source)
	records = compare_trucks(parsed["trucks"])
	invalid_rows = [
		{"row_number": row.row_number, "value": row.value, "reason": row.reason}
		for row in parsed["invalid_rows"]
	]
	return {
		"summary": build_summary(parsed["source_rows"], parsed["trucks"], records, parsed["invalid_rows"]),
		"records": records,
		"invalid_rows": invalid_rows,
	}


def import_new_owned_trucks(source: str | Path) -> dict[str, Any]:
	result = dry_run_owned_truck_import(source)
	records = result["records"]
	summary = result["summary"]
	created = 0
	failed_rows = []

	if summary["conflict"] or summary["ownership_conflict"]:
		result["summary"] = build_import_summary(
			summary,
			created=0,
			skipped=sum(1 for row in records if row["status"] != "NEW"),
			failed=summary["invalid"] + summary["conflict"] + summary["ownership_conflict"],
			created_truck_types=0,
		)
		return result

	created_truck_types = ensure_required_truck_types(records)

	for row in records:
		if row["status"] != "NEW":
			continue

		try:
			doc = frappe.new_doc("Truck")
			doc.truck_number = row["normalized_plate"]
			doc.license_plate = row["normalized_plate"]
			doc.vehicle_type = row["source_truck_type"]
			doc.ownership_type = OWNERSHIP_TYPE
			doc.status = DEFAULT_STATUS
			doc.disabled = 0
			doc.insert()
			row["status"] = "CREATED"
			row["created_name"] = doc.name
			created += 1
		except Exception as exc:
			frappe.db.rollback()
			row["status"] = "FAILED"
			row["error"] = str(exc)
			failed_rows.append(row)
		else:
			frappe.db.commit()

	failed = len(failed_rows) + len(result["invalid_rows"])
	skipped = sum(
		1
		for row in records
		if row["status"] in {"EXISTS", "UPDATE_TYPE", "OWNERSHIP_CONFLICT", "DUPLICATE_SOURCE", "CONFLICT"}
	)
	result["summary"] = build_import_summary(
		summary,
		created=created,
		skipped=skipped,
		failed=failed,
		created_truck_types=created_truck_types,
	)
	result["failed_rows"] = failed_rows
	return result


def ensure_required_truck_types(records: list[dict[str, Any]]) -> int:
	required_types = {
		row["source_truck_type"]
		for row in records
		if row["status"] == "NEW" and row["source_truck_type"] in SUPPORTED_TRUCK_TYPES
	}
	created = 0
	for truck_type in sorted(required_types):
		if frappe.db.exists("Truck Type", truck_type):
			continue
		doc = frappe.new_doc("Truck Type")
		doc.truck_type = truck_type
		doc.insert()
		created += 1
	frappe.db.commit()
	return created


def build_import_summary(
	summary: dict[str, Any],
	*,
	created: int,
	skipped: int,
	failed: int,
	created_truck_types: int,
) -> dict[str, Any]:
	import_summary = summary.copy()
	import_summary.update(
		{
			"created": created,
			"created_truck_types": created_truck_types,
			"skipped": skipped,
			"failed": failed,
			"final_count": frappe.db.count("Truck"),
		}
	)
	return import_summary
